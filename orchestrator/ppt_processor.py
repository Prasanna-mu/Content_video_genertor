import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Emu
import json

from models.database import SessionLocal
from models.models import PPTJob, PPTJobStatus, Slide, SlideStatus
from config.settings import settings
from utils.json_utils import save_json
from utils.logging import get_logger
from utils.filesystem import ensure_dir


logger = get_logger("ppt_processor")


class PPTProcessor:
    def __init__(self, session_id: str):
        self.session_id = session_id
        self.db = SessionLocal()
        self.parse_output_dir = settings.PROCESSING_DIR / "parseppt" / session_id
        self.parse_output_dir.mkdir(parents=True, exist_ok=True)

    def process_job(self, job_id: str) -> bool:
        job = self.db.query(PPTJob).filter(PPTJob.id == job_id).first()
        if not job:
            logger.error(f"[{self.session_id}] PPT job not found: {job_id}")
            return False

        job.status = PPTJobStatus.IN_PROGRESS
        job.started_at = datetime.utcnow()
        self.db.commit()

        ppt_path = settings.INPUT_PPT_DIR / job.ppt_filename
        if not ppt_path.exists():
            error_msg = f"PPT file not found: {ppt_path}"
            logger.error(f"[{self.session_id}] {error_msg}")
            job.status = PPTJobStatus.FAILED
            job.error = error_msg
            job.completed_at = datetime.utcnow()
            self.db.commit()
            return False

        try:
            parsed_data = self._parse_ppt(ppt_path)
            self._save_parsed_json(job, parsed_data)
            self._create_slide_records(job, parsed_data)

            job.status = PPTJobStatus.COMPLETED
            job.completed_at = datetime.utcnow()
            self.db.commit()
            logger.info(f"[{self.session_id}] Parsed PPT: {job.ppt_filename} ({parsed_data['ppt']['slide_count']} slides)")
            return True

        except Exception as e:
            error_msg = f"Failed to parse PPT {job.ppt_filename}: {str(e)}"
            logger.error(f"[{self.session_id}] {error_msg}")
            job.status = PPTJobStatus.FAILED
            job.error = error_msg
            job.completed_at = datetime.utcnow()
            self.db.commit()
            return False

    def _parse_ppt(self, ppt_path: Path) -> Dict[str, Any]:
        prs = Presentation(ppt_path)

        slides_data = []
        for i, slide in enumerate(prs.slides):
            slide_data = self._parse_slide(slide, i + 1)
            slides_data.append(slide_data)

        return {
            "ppt": {
                "filename": ppt_path.name,
                "slide_count": len(slides_data),
            },
            "slides": slides_data,
        }

    def _parse_slide(self, slide, slide_number: int) -> Dict[str, Any]:
        title = ""
        text_boxes = []
        bullet_points = []
        notes = ""
        tables = []
        images = []
        shapes = []

        for shape in slide.shapes:
            shape_info = {
                "shape_id": shape.shape_id,
                "name": shape.name,
                "left": shape.left,
                "top": shape.top,
                "width": shape.width,
                "height": shape.height,
                "shape_type": str(shape.shape_type),
            }

            if shape.has_text_frame:
                text_content = self._extract_text_frame(shape.text_frame)
                if text_content:
                    if self._is_title_shape(shape, slide):
                        title = text_content
                    else:
                        text_boxes.append({
                            "shape_id": shape.shape_id,
                            "text": text_content,
                            "paragraphs": self._extract_paragraphs(shape.text_frame),
                        })

            if shape.has_table:
                table_data = self._extract_table(shape.table)
                tables.append({
                    "shape_id": shape.shape_id,
                    "rows": table_data["rows"],
                    "cols": table_data["cols"],
                    "data": table_data["data"],
                })

            if self._is_image_shape(shape):
                image_info = self._extract_image_info(shape)
                if image_info:
                    images.append(image_info)

            shapes.append(shape_info)

        if slide.has_notes_slide:
            notes = self._extract_notes(slide.notes_slide)

        bullet_points = self._extract_bullets(text_boxes)

        return {
            "slide_number": slide_number,
            "title": title,
            "text_boxes": text_boxes,
            "bullet_points": bullet_points,
            "notes": notes,
            "tables": tables,
            "images": images,
            "shapes_count": len(shapes),
        }

    def _extract_text_frame(self, text_frame) -> str:
        paragraphs = []
        for para in text_frame.paragraphs:
            text = "".join(run.text for run in para.runs).strip()
            if text:
                paragraphs.append(text)
        return "\n".join(paragraphs)

    def _extract_paragraphs(self, text_frame) -> List[Dict[str, Any]]:
        paragraphs = []
        for para in text_frame.paragraphs:
            runs_data = []
            for run in para.runs:
                runs_data.append({
                    "text": run.text,
                    "bold": run.font.bold,
                    "italic": run.font.italic,
                    "font_size": run.font.size.pt if run.font.size else None,
                    "font_name": run.font.name,
                })
            paragraphs.append({
                "text": "".join(run.text for run in para.runs).strip(),
                "level": para.level,
                "runs": runs_data,
            })
        return paragraphs

    def _is_title_shape(self, shape, slide) -> bool:
        if shape.shape_type == 17:
            return True
        if hasattr(shape, 'placeholder_format') and shape.placeholder_format is not None:
            try:
                if shape.placeholder_format.type == 1:
                    return True
            except Exception:
                pass
        if hasattr(slide, 'shapes') and hasattr(slide.shapes, 'title') and slide.shapes.title == shape:
            return True
        return False

    def _extract_bullets(self, text_boxes: List[Dict]) -> List[str]:
        bullets = []
        for box in text_boxes:
            for para in box.get("paragraphs", []):
                if para.get("level", 0) > 0 or para.get("text", "").startswith(("•", "-", "*", "→", "▸")):
                    bullets.append(para["text"])
        return bullets

    def _extract_table(self, table) -> Dict[str, Any]:
        data = []
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip())
            data.append(row_data)
        return {
            "rows": len(table.rows),
            "cols": len(table.columns),
            "data": data,
        }

    def _is_image_shape(self, shape) -> bool:
        return shape.shape_type == 13 or (
            hasattr(shape, "image") and shape.image is not None
        )

    def _extract_image_info(self, shape) -> Optional[Dict[str, Any]]:
        try:
            if not hasattr(shape, "image") or shape.image is None:
                return None
            img = shape.image
            return {
                "shape_id": shape.shape_id,
                "content_type": img.content_type,
                "size_bytes": len(img.blob),
                "width": shape.width,
                "height": shape.height,
            }
        except Exception:
            return None

    def _extract_notes(self, notes_slide) -> str:
        text_parts = []
        for shape in notes_slide.shapes:
            if shape.has_text_frame:
                text = self._extract_text_frame(shape.text_frame)
                if text:
                    text_parts.append(text)
        return "\n".join(text_parts)

    def _save_parsed_json(self, job: PPTJob, parsed_data: Dict[str, Any]) -> Path:
        output_filename = Path(job.ppt_filename).stem + ".json"
        output_path = self.parse_output_dir / output_filename
        save_json(output_path, parsed_data)
        return output_path

    def _create_slide_records(self, job: PPTJob, parsed_data: Dict[str, Any]):
        for slide_data in parsed_data["slides"]:
            slide_id = f"{job.id}_slide_{slide_data['slide_number']}"
            slide = Slide(
                id=slide_id,
                ppt_job_id=job.id,
                slide_number=slide_data["slide_number"],
                title=slide_data["title"] or None,
                raw_text=json.dumps(slide_data["text_boxes"], ensure_ascii=False),
                notes=slide_data["notes"] or None,
                status=SlideStatus.PARSED,
                parsed=1,
            )
            self.db.add(slide)
        self.db.commit()

    def close(self):
        self.db.close()


def discover_ppt_files() -> List[str]:
    ppt_dir = settings.INPUT_PPT_DIR
    ensure_dir(ppt_dir)
    
    pptx_files = list(ppt_dir.glob("*.pptx"))
    pptx_files += list(ppt_dir.glob("*.ppt"))
    
    filenames = [f.name for f in pptx_files]
    logger.info(f"Discovered {len(filenames)} PPT file(s) in {ppt_dir}: {filenames}")
    return filenames


def create_ppt_jobs(session_id: str, ppt_filenames: List[str]) -> List[PPTJob]:
    db = SessionLocal()
    created_jobs = []

    for filename in ppt_filenames:
        # Generate session-specific job ID to avoid conflicts across sessions
        base_name = Path(filename).stem
        job_id = f"{session_id}_{base_name}"

        # Check if job already exists (for resumability)
        existing = db.query(PPTJob).filter(PPTJob.id == job_id).first()
        if existing:
            # Verify the file still exists
            ppt_path = settings.INPUT_PPT_DIR / existing.ppt_filename
            if ppt_path.exists():
                created_jobs.append(existing)
                continue
            else:
                logger.warning(f"[{session_id}] PPT file missing for existing job {existing.id}: {existing.ppt_filename}")
                existing.status = PPTJobStatus.FAILED
                existing.error = f"Source file no longer exists: {existing.ppt_filename}"
                db.commit()

        ppt_job = PPTJob(
            id=job_id,
            session_id=session_id,
            ppt_filename=filename,
            status=PPTJobStatus.PENDING,
        )
        db.add(ppt_job)
        created_jobs.append(ppt_job)

    db.commit()
    for job in created_jobs:
        db.refresh(job)
    db.close()

    return created_jobs